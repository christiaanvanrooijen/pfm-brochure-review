"""
Dress the translation export as a review workbook for native speakers.

    node --experimental-strip-types scripts/translation-review-export.mjs > /tmp/t.json
    python3 scripts/translation-review-workbook.py /tmp/t.json docs/content/PFM-translation-review-FR-DE.xlsx

Tabs: Read me · French review · German review · Not translated.
The yellow columns are the reviewer's; everything else is generated and should
not be edited. Re-running replaces the file, so collect reviews before that.
"""

from __future__ import annotations

import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF4CC")
HEAD = PatternFill("solid", fgColor="0C111D")
HIGH = PatternFill("solid", fgColor="F4EBFF")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

AREA_ORDER = [
    "Cover (Unified PFM Intro)", "Segment picker", "Customer cases",
    "Drawer · limitations", "Drawer · how it works (method copy)",
    "Drawer · scene and segment purpose", "Drawer · measurement illustrations",
    "Interface", "Segment covers and first questions", "Scenes",
    "Journey · Retail", "Journey · Retail Park", "Journey · Outlet Centre", "Journey · Drive-Thru",
    "Capabilities", "Technology · what it does", "Technology · limitation (fallback)",
    "Drawer · example videos", "Privacy statement",
    "Requirements · data roles", "Requirements · installation essentials",
    "Requirements · technical detail", "Requirements · inputs",
]


def order(row):
    area = AREA_ORDER.index(row["area"]) if row["area"] in AREA_ORDER else len(AREA_ORDER)
    return (0 if row["priority"].startswith("High") else 1, area)


def header(ws, cols):
    for i, (title, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def review_sheet(wb, title, lang, key, rows):
    ws = wb.create_sheet(title)
    cols = [
        ("#", 6), ("Priority", 14), ("Area", 24), ("Where it appears", 30),
        ("English (source)", 55), (f"{lang} (current)", 55),
        ("Verdict", 12), (f"Suggested {lang}", 50), ("Reviewer comment", 34),
    ]
    header(ws, cols)
    dv = DataValidation(type="list", formula1='"OK,Change,Question"', allow_blank=True)
    ws.add_data_validation(dv)
    for n, row in enumerate(rows, start=2):
        values = [n - 1, row["priority"], row["area"], row["where"], row["en"], row[key], None, None, None]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=n, column=col, value=value)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            c.border = BORDER
            if col >= 7:
                c.fill = YELLOW
            elif row["priority"].startswith("High") and col <= 2:
                c.fill = HIGH
        dv.add(f"G{n}")
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
    return ws


def main(src: str, out: str) -> None:
    rows = json.load(open(src, encoding="utf-8"))
    fr_rows = sorted([r for r in rows if not r["frMissing"]], key=order)
    de_rows = sorted([r for r in rows if not r["deMissing"]], key=order)
    missing = sorted([r for r in rows if r["frMissing"] or r["deMissing"]], key=order)

    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    lines = [
        ("PFM Commercial Experience — French and German translation review", None),
        ("Generated", date.today().isoformat()),
        ("What this is", "Every French and German text the digital brochure can show, beside its English source, read from the running app's own translation layer. Identical texts are listed once, with every place they appear."),
        ("Who reviews what", "A native French speaker reviews the 'French review' tab; a native German speaker the 'German review' tab. The 'Not translated' tab is for the product lead, not for translators."),
        ("How to review", "Read the English, then the current translation. In the yellow columns: set Verdict to OK, Change or Question. For Change, write your version in 'Suggested …'. Use 'Reviewer comment' for anything else. Leave the white columns as they are."),
        ("Start with", "Rows marked 'High — new this week' (purple). They were written or rewritten on 2026-09-27/28 and have not been read by a native speaker yet. Everything else predates them."),
        ("Keep the meaning, not just the words", "The English is precise on purpose. Keep these distinctions: a visit is not a passer-by; movement events are not unique visitors; 'illustrative' means example data, not a result; nothing is presented as measured when it is derived or estimated. Do not add figures, promises or claims the English does not make. Brand and product names (PFM, People Flow Management, Madaq, Future Stores London, Advantage) stay as they are."),
        ("Register", "Professional B2B, addressed to retail and property decision-makers. French: vouvoiement. German: Sie."),
        ("Example of a filled row", "Verdict: Change · Suggested French: « Libérez le potentiel de chaque site. » · Reviewer comment: « lieu » sounds too generic for a retail property audience."),
        ("Progress — French", None),
        ("Progress — German", None),
        ("Counts", f"{len(fr_rows)} French rows · {len(de_rows)} German rows · {len(missing)} texts still shown in English"),
        ("Source", "scripts/translation-review-export.mjs and scripts/translation-review-workbook.py, from the typed runtime. Regenerating replaces this file — collect reviews first."),
    ]
    for i, (label, text) in enumerate(lines, start=1):
        a = ws.cell(row=i, column=1, value=label)
        a.font = Font(name=FONT, bold=True, size=14 if i == 1 else 10)
        a.alignment = WRAP
        if text is not None:
            b = ws.cell(row=i, column=2, value=text)
            b.font = Font(name=FONT, size=10)
            b.alignment = WRAP
    # Live progress, counted from the review tabs.
    fr_n, de_n = len(fr_rows) + 1, len(de_rows) + 1
    ws["B10"] = f"=COUNTA('French review'!G2:G{fr_n})&\" of {len(fr_rows)} reviewed · \"&COUNTIF('French review'!G2:G{fr_n},\"Change\")&\" to change · \"&COUNTIF('French review'!G2:G{fr_n},\"Question\")&\" questions\""
    ws["B11"] = f"=COUNTA('German review'!G2:G{de_n})&\" of {len(de_rows)} reviewed · \"&COUNTIF('German review'!G2:G{de_n},\"Change\")&\" to change · \"&COUNTIF('German review'!G2:G{de_n},\"Question\")&\" questions\""
    for ref in ("B10", "B11"):
        ws[ref].font = Font(name=FONT, size=10)
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=False)

    review_sheet(wb, "French review", "French", "fr", fr_rows)
    review_sheet(wb, "German review", "German", "de", de_rows)

    ws = wb.create_sheet("Not translated")
    header(ws, [("#", 6), ("Area", 24), ("Where it appears", 34), ("English (shown instead)", 60),
                ("French", 16), ("German", 16), ("Product lead: translate or leave?", 34)])
    for n, row in enumerate(missing, start=2):
        values = [n - 1, row["area"], row["where"], row["en"],
                  "English shown" if row["frMissing"] else "translated",
                  "English shown" if row["deMissing"] else "translated", None]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=n, column=col, value=value)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            c.border = BORDER
            if col == 7:
                c.fill = YELLOW
    ws.auto_filter.ref = f"A1:G{len(missing) + 1}"

    # The two progress formulas have no cached value until a spreadsheet
    # application computes them; ask it to on open.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    print(f"{out}: {len(fr_rows)} FR, {len(de_rows)} DE, {len(missing)} not translated")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
